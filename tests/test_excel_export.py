"""
Tests for excel_export.py. The structural/shape tests always run. The
`TestExcelFormulasRecalculateCorrectly` class headlessly recalculates the
generated workbook with the `formulas` library and cross-checks every
derived figure against dcf_engine.py's own Python calculation for the same
inputs — this is the "verify the exported workbook's formulas actually
compute correctly" check from the project brief. Skipped automatically if
`formulas` isn't installed (it's a dev-only verification dependency, not a
runtime one — see requirements-dev.txt).
"""
import pandas as pd
import pytest

import dcf_engine as eng
import excel_export as xl

try:
    import formulas
    HAS_FORMULAS = True
except ImportError:
    HAS_FORMULAS = False


def _build_payload_and_python_results():
    """One consistent set of synthetic inputs, run through both dcf_engine.py and excel_export.py, for comparison."""
    financials = pd.DataFrame({
        "Total Revenue": [1000.0, 900.0, 820.0, 750.0],
        "EBIT": [150.0, 130.0, 115.0, 100.0],
        "Tax Provision": [36.0, 31.0, 27.0, 24.0],
        "Pretax Income": [140.0, 122.0, 108.0, 94.0],
    }, index=pd.to_datetime(["2025-06-30", "2024-06-30", "2023-06-30", "2022-06-30"])).T
    cashflow = pd.DataFrame({
        "Depreciation And Amortization": [40.0, 36.0, 33.0, 30.0],
        "Capital Expenditure": [-60.0, -54.0, -50.0, -45.0],
        "Change In Working Capital": [-10.0, -8.0, -5.0, -4.0],
    }, index=pd.to_datetime(["2025-06-30", "2024-06-30", "2023-06-30", "2022-06-30"])).T
    balance_sheet = pd.DataFrame({
        "Total Debt": [300.0, 310.0, 320.0, 330.0],
        "Cash And Cash Equivalents": [150.0, 130.0, 110.0, 90.0],
    }, index=pd.to_datetime(["2025-06-30", "2024-06-30", "2023-06-30", "2022-06-30"])).T

    hist = eng.extract_historical_financials(financials, cashflow, balance_sheet)
    snap = eng.extract_company_snapshot({"currentPrice": 12.5, "marketCap": 5000.0, "beta": 1.1, "sharesOutstanding": 400.0}, balance_sheet)
    assumptions = eng.build_default_forecast_assumptions(hist, n_years=5, terminal_growth_pct=2.5)
    wacc_res = eng.compute_wacc(0.043, snap.beta, 0.055, interest_expense=15.0, total_debt=snap.total_debt, market_cap=snap.market_cap, tax_rate=hist.tax_rate_avg)
    proj = eng.compute_fcf_projection(assumptions, hist.latest_revenue, hist.tax_rate_avg, wacc_res.wacc)
    tv = eng.compute_terminal_value(
        fcf_final=proj["fcf"].iloc[-1], ebitda_final=proj["ebitda"].iloc[-1], n_years=len(proj),
        wacc=wacc_res.wacc, terminal_growth=0.025, exit_multiple=8.0, perpetuity_weight=0.5,
    )
    bridge = eng.compute_valuation_bridge(proj["pv_fcf"].sum(), tv.pv_terminal_value, snap.total_debt, snap.cash, snap.shares_outstanding, snap.current_price)

    payload = xl.ExportPayload(
        ticker="TEST.AX", company_name="Test Co", current_price=snap.current_price,
        tax_rate=hist.tax_rate_avg,
        risk_free_rate=0.043, equity_risk_premium=0.055, beta=snap.beta,
        interest_expense=15.0, total_debt=snap.total_debt, market_cap=snap.market_cap,
        base_revenue=hist.latest_revenue, assumptions=assumptions,
        terminal_growth=0.025, exit_multiple=8.0, perpetuity_weight=0.5,
        latest_historical_ebitda=hist.latest_ebitda,
        cash=snap.cash, net_debt_override=None, shares_outstanding=snap.shares_outstanding, shares_override=None,
    )
    python_results = {
        "wacc": wacc_res.wacc,
        "sum_pv_fcf": float(proj["pv_fcf"].sum()),
        "tv_perpetuity": tv.tv_perpetuity,
        "tv_exit": tv.tv_exit_multiple,
        "blended_tv": tv.blended_tv,
        "pv_tv": tv.pv_terminal_value,
        "implied_exit_multiple": tv.implied_exit_multiple_from_perpetuity,
        "implied_perp_growth": tv.implied_perpetuity_growth_from_exit,
        "enterprise_value": bridge.enterprise_value,
        "equity_value": bridge.equity_value,
        "value_per_share": bridge.value_per_share,
        "upside_downside": bridge.upside_downside,
    }
    return payload, python_results


# ---------------------------------------------------------------------------
# Structural tests — always run
# ---------------------------------------------------------------------------

class TestWorkbookStructure:
    def test_builds_without_error(self):
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        assert wb is not None

    def test_has_expected_sheets(self):
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        assert wb.sheetnames == ["Assumptions", "WACC", "FCF Projection", "Terminal Value", "Valuation Summary"]

    def test_fcf_projection_has_one_column_per_forecast_year(self):
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        ws = wb["FCF Projection"]
        # year row (row 2): columns C.. should hold 1..n_years
        n_years = len(payload.assumptions)
        years = [ws.cell(row=2, column=c).value for c in range(3, 3 + n_years)]
        assert years == list(range(1, n_years + 1))

    def test_fcf_projection_column_count_matches_different_n_years(self):
        payload, _ = _build_payload_and_python_results()
        payload.assumptions = payload.assumptions.iloc[:3].reset_index(drop=True)  # only 3 forecast years
        wb = xl.build_workbook(payload)
        ws = wb["FCF Projection"]
        years = [ws.cell(row=2, column=c).value for c in range(3, 6)]
        assert years == [1, 2, 3]
        assert ws.cell(row=2, column=6).value is None  # no 4th year column

    def test_input_cells_are_styled_blue_on_yellow(self):
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        cell = wb["WACC"]["B3"]  # risk-free rate input
        assert cell.font.color.rgb.endswith("0000FF")
        assert cell.fill.fgColor.rgb.endswith("FFFF99")

    def test_formula_cells_are_not_input_styled(self):
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        cell = wb["WACC"]["B16"]  # WACC formula result
        assert str(cell.value).startswith("=")
        assert cell.font.color.rgb.endswith("000000")

    def test_no_cell_value_is_a_fake_formula_string(self):
        """Regression guard: a plain annotation string that happens to start with '=' gets
        silently treated as a real formula by Excel instead of displayed as text."""
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("=") and "->" in cell.value:
                        pytest.fail(f"{ws.title}!{cell.coordinate} looks like an annotation string mistakenly starting with '=': {cell.value!r}")

    def test_export_to_bytes_round_trips(self):
        from openpyxl import load_workbook
        import io
        payload, _ = _build_payload_and_python_results()
        data = xl.export_to_bytes(payload)
        assert isinstance(data, bytes) and len(data) > 0
        wb = load_workbook(io.BytesIO(data))
        assert "Valuation Summary" in wb.sheetnames


# ---------------------------------------------------------------------------
# Formula correctness — recalculated headlessly, cross-checked vs Python
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def recalculated():
    if not HAS_FORMULAS:
        pytest.skip("'formulas' package not installed — see requirements-dev.txt")
    payload, python_results = _build_payload_and_python_results()
    wb = xl.build_workbook(payload)
    import tempfile, os
    tmp_dir = tempfile.mkdtemp()
    path = os.path.join(tmp_dir, "test_model.xlsx")
    wb.save(path)

    xl_model = formulas.ExcelModel().loads(path).finish()
    solution = xl_model.calculate()
    basename = os.path.basename(path)

    def get(sheet: str, cell: str):
        key = f"'[{basename}]{sheet.upper()}'!{cell}"
        return solution[key].value[0][0]

    return get, python_results


@pytest.mark.skipif(not HAS_FORMULAS, reason="'formulas' package not installed — see requirements-dev.txt")
class TestExcelFormulasRecalculateCorrectly:
    def test_no_error_cells_anywhere(self, recalculated):
        get, _ = recalculated
        # re-derive the raw solution to scan every cell, not just the ones we assert on below
        payload, _ = _build_payload_and_python_results()
        wb = xl.build_workbook(payload)
        import tempfile, os
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "scan.xlsx")
            wb.save(path)
            xl_model = formulas.ExcelModel().loads(path).finish()
            solution = xl_model.calculate()
            error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!")
            errors = [(k, v.value) for k, v in solution.items() if any(tok in str(v.value) for tok in error_tokens)]
        assert errors == [], f"Excel error values found: {errors}"

    def test_wacc_matches_python(self, recalculated):
        get, py = recalculated
        assert get("WACC", "B16") == pytest.approx(py["wacc"], rel=1e-6)

    def test_sum_pv_fcf_matches_python(self, recalculated):
        get, py = recalculated
        assert get("FCF Projection", "B18") == pytest.approx(py["sum_pv_fcf"], rel=1e-6)

    def test_terminal_value_matches_python(self, recalculated):
        get, py = recalculated
        assert get("Terminal Value", "B7") == pytest.approx(py["tv_perpetuity"], rel=1e-6)
        assert get("Terminal Value", "B8") == pytest.approx(py["tv_exit"], rel=1e-6)
        assert get("Terminal Value", "B9") == pytest.approx(py["blended_tv"], rel=1e-6)
        assert get("Terminal Value", "B10") == pytest.approx(py["pv_tv"], rel=1e-6)

    def test_terminal_value_cross_checks_match_python(self, recalculated):
        get, py = recalculated
        assert get("Terminal Value", "B13") == pytest.approx(py["implied_exit_multiple"], rel=1e-6)
        assert get("Terminal Value", "B14") == pytest.approx(py["implied_perp_growth"], rel=1e-6)

    def test_valuation_summary_matches_python(self, recalculated):
        get, py = recalculated
        assert get("Valuation Summary", "B5") == pytest.approx(py["enterprise_value"], rel=1e-6)
        assert get("Valuation Summary", "B7") == pytest.approx(py["equity_value"], rel=1e-6)
        assert get("Valuation Summary", "B9") == pytest.approx(py["value_per_share"], rel=1e-6)
        assert get("Valuation Summary", "B11") == pytest.approx(py["upside_downside"], rel=1e-6)

    def test_net_debt_override_flows_through_excel_formula(self, tmp_path):
        """Net Debt override uses IF(Assumptions!B11<>"",...) — verify the override branch actually fires in Excel, not just the default branch."""
        payload, _ = _build_payload_and_python_results()
        payload.net_debt_override = 999.0
        wb = xl.build_workbook(payload)
        path = tmp_path / "override_test.xlsx"
        wb.save(path)
        xl_model = formulas.ExcelModel().loads(str(path)).finish()
        solution = xl_model.calculate()
        net_debt = solution[f"'[{path.name}]VALUATION SUMMARY'!B6"].value[0][0]
        assert net_debt == pytest.approx(999.0)
