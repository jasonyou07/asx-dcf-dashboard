"""
Builds a formula-driven .xlsx mirror of the dashboard using openpyxl.

Every number the dashboard used at export time becomes a BLUE/YELLOW input
cell; every derived number is a real Excel formula referencing those cells
by address — not a pasted-in Python result. Editing an input cell in Excel
recalculates the whole chain natively.

Sheet layout: Assumptions & Inputs -> WACC Calculation -> FCF Projection
(columns sized to however many forecast years were actually set) ->
Terminal Value -> Valuation Summary.
"""

from dataclasses import dataclass

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Styling — standard financial-model convention
# ---------------------------------------------------------------------------

INPUT_FONT = Font(color="0000FF", bold=False)                 # blue font
INPUT_FILL = PatternFill("solid", fgColor="FFFF99")            # yellow fill
FORMULA_FONT = Font(color="000000", bold=False)                 # black font
HEADER_FONT = Font(bold=True, size=13, color="1F3864")
SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="1F3864")
LABEL_FONT = Font(bold=False, color="000000")
THIN_BORDER = Border(bottom=Side(style="thin", color="CCCCCC"))

PCT_FORMAT = "0.00%"
CURRENCY_FORMAT = "#,##0.00"
MULTIPLE_FORMAT = "0.00\"x\""
NUMBER_FORMAT = "#,##0.00"


def _input_cell(ws: Worksheet, coord: str, value, number_format: str = NUMBER_FORMAT):
    cell = ws[coord]
    cell.value = value
    cell.font = INPUT_FONT
    cell.fill = INPUT_FILL
    cell.number_format = number_format
    return cell


def _formula_cell(ws: Worksheet, coord: str, formula: str, number_format: str = NUMBER_FORMAT):
    cell = ws[coord]
    cell.value = formula
    cell.font = FORMULA_FONT
    cell.number_format = number_format
    return cell


def _label(ws: Worksheet, coord: str, text: str, bold: bool = False):
    cell = ws[coord]
    cell.value = text
    cell.font = Font(bold=bold)
    return cell


def _section_header(ws: Worksheet, row: int, text: str, span: int = 4):
    for col in range(1, span + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = SECTION_FILL
        cell.font = SECTION_FONT
    ws.cell(row=row, column=1).value = text


# ---------------------------------------------------------------------------
# Export payload — everything the workbook needs, gathered from the
# dashboard's live state at export time. Plain dataclass (not the dashboard's
# own session state) so this module stays independently testable.
# ---------------------------------------------------------------------------

@dataclass
class ExportPayload:
    ticker: str
    company_name: str
    current_price: float | None

    # shared
    tax_rate: float  # fraction

    # WACC inputs
    risk_free_rate: float
    equity_risk_premium: float
    beta: float
    interest_expense: float | None
    total_debt: float
    market_cap: float

    # FCF forecast
    base_revenue: float
    assumptions: pd.DataFrame  # columns: year, revenue_growth_pct, ebit_margin_pct, capex_pct_of_revenue, da_pct_of_revenue, nwc_pct_of_revenue

    # terminal value
    terminal_growth: float   # fraction
    exit_multiple: float
    perpetuity_weight: float  # fraction
    latest_historical_ebitda: float | None

    # bridge
    cash: float
    net_debt_override: float | None
    shares_outstanding: float
    shares_override: float | None


def build_workbook(payload: ExportPayload) -> Workbook:
    wb = Workbook()
    ws_assump = wb.active
    ws_assump.title = "Assumptions"
    ws_wacc = wb.create_sheet("WACC")
    ws_fcf = wb.create_sheet("FCF Projection")
    ws_tv = wb.create_sheet("Terminal Value")
    ws_sum = wb.create_sheet("Valuation Summary")

    n_years = len(payload.assumptions)

    # ------------------------------------------------------------------
    # Assumptions & Inputs
    # ------------------------------------------------------------------
    ws_assump["A1"] = f"{payload.company_name or payload.ticker} ({payload.ticker}) — DCF Assumptions & Inputs"
    ws_assump["A1"].font = HEADER_FONT
    ws_assump.column_dimensions["A"].width = 32
    ws_assump.column_dimensions["B"].width = 16

    _label(ws_assump, "A3", "Current Price")
    _input_cell(ws_assump, "B3", payload.current_price or 0.0, CURRENCY_FORMAT)
    _label(ws_assump, "A4", "Shared Tax Rate")
    _input_cell(ws_assump, "B4", payload.tax_rate, PCT_FORMAT)
    ws_assump["C4"] = "(used in both FCF tax shield and after-tax cost of debt)"

    _label(ws_assump, "A6", "Base-year Revenue")
    _input_cell(ws_assump, "B6", payload.base_revenue, NUMBER_FORMAT)
    _label(ws_assump, "A7", "Market Cap")
    _input_cell(ws_assump, "B7", payload.market_cap, NUMBER_FORMAT)
    _label(ws_assump, "A8", "Total Debt")
    _input_cell(ws_assump, "B8", payload.total_debt, NUMBER_FORMAT)
    _label(ws_assump, "A9", "Cash")
    _input_cell(ws_assump, "B9", payload.cash, NUMBER_FORMAT)
    _label(ws_assump, "A10", "Shares Outstanding")
    _input_cell(ws_assump, "B10", payload.shares_override if payload.shares_override is not None else payload.shares_outstanding, NUMBER_FORMAT)
    _label(ws_assump, "A11", "Net Debt Override (blank cell B8-B9 used if empty)")
    if payload.net_debt_override is not None:
        _input_cell(ws_assump, "B11", payload.net_debt_override, NUMBER_FORMAT)
    _label(ws_assump, "A12", "Latest Historical EBITDA (actual, not forecast)")
    _input_cell(ws_assump, "B12", payload.latest_historical_ebitda or 0.0, NUMBER_FORMAT)

    # ------------------------------------------------------------------
    # WACC Calculation
    # ------------------------------------------------------------------
    ws_wacc["A1"] = "WACC Calculation"
    ws_wacc["A1"].font = HEADER_FONT
    ws_wacc.column_dimensions["A"].width = 30
    ws_wacc.column_dimensions["B"].width = 16
    ws_wacc.column_dimensions["C"].width = 40

    _label(ws_wacc, "A3", "Risk-Free Rate")
    _input_cell(ws_wacc, "B3", payload.risk_free_rate, PCT_FORMAT)
    _label(ws_wacc, "A4", "Equity Risk Premium")
    _input_cell(ws_wacc, "B4", payload.equity_risk_premium, PCT_FORMAT)
    _label(ws_wacc, "A5", "Beta")
    _input_cell(ws_wacc, "B5", payload.beta, "0.00")
    _label(ws_wacc, "A6", "Interest Expense")
    _input_cell(ws_wacc, "B6", abs(payload.interest_expense) if payload.interest_expense is not None else 0.0, NUMBER_FORMAT)
    _label(ws_wacc, "A7", "Total Debt")
    _formula_cell(ws_wacc, "B7", "=Assumptions!B8", NUMBER_FORMAT)
    _label(ws_wacc, "A8", "Market Cap")
    _formula_cell(ws_wacc, "B8", "=Assumptions!B7", NUMBER_FORMAT)
    _label(ws_wacc, "A9", "Tax Rate")
    _formula_cell(ws_wacc, "B9", "=Assumptions!B4", PCT_FORMAT)

    _label(ws_wacc, "A11", "Cost of Equity = Risk-Free + Beta x ERP", bold=True)
    _formula_cell(ws_wacc, "B11", "=B3+B5*B4", PCT_FORMAT)
    ws_wacc["C11"] = "B3+B5*B4 -> Risk-Free Rate + Beta x Equity Risk Premium"

    _label(ws_wacc, "A12", "Cost of Debt = Interest Expense / Total Debt", bold=True)
    _formula_cell(ws_wacc, "B12", '=IF(B7>0,B6/B7,B3+0.02)', PCT_FORMAT)
    ws_wacc["C12"] = "falls back to Risk-Free + 2% if Total Debt is 0"

    _label(ws_wacc, "A13", "Weight of Equity = Market Cap / (Market Cap + Debt)", bold=True)
    _formula_cell(ws_wacc, "B13", "=IF((B8+B7)>0,B8/(B8+B7),1)", PCT_FORMAT)
    _label(ws_wacc, "A14", "Weight of Debt = 1 - Weight of Equity", bold=True)
    _formula_cell(ws_wacc, "B14", "=1-B13", PCT_FORMAT)

    _label(ws_wacc, "A16", "WACC = WeightEq x CostEq + WeightDebt x CostDebt x (1-Tax)", bold=True)
    _formula_cell(ws_wacc, "B16", "=B13*B11+B14*B12*(1-B9)", PCT_FORMAT)
    ws_wacc["B16"].font = Font(bold=True, color="000000")

    # ------------------------------------------------------------------
    # FCF Projection (dynamic columns, one per forecast year)
    # ------------------------------------------------------------------
    ws_fcf["A1"] = "FCF Projection"
    ws_fcf["A1"].font = HEADER_FONT
    ws_fcf.column_dimensions["A"].width = 28

    row_labels = [
        ("Year", None),
        ("Revenue Growth %", "pct_input"),
        ("Revenue", "formula"),
        ("EBIT Margin %", "pct_input"),
        ("EBIT", "formula"),
        ("D&A % of Revenue", "pct_input"),
        ("D&A", "formula"),
        ("CapEx % of Revenue", "pct_input"),
        ("CapEx", "formula"),
        ("ΔNWC % of Revenue", "pct_input"),
        ("ΔNWC", "formula"),
        ("EBITDA (EBIT + D&A)", "formula"),
        ("FCF = EBIT x (1-Tax) + D&A - CapEx - ΔNWC", "formula"),
        ("Discount Factor = 1/(1+WACC)^t", "formula"),
        ("PV of FCF", "formula"),
    ]
    for i, (label, _) in enumerate(row_labels, start=2):
        _label(ws_fcf, f"A{i}", label, bold=(label in ("FCF = EBIT x (1-Tax) + D&A - CapEx - ΔNWC", "PV of FCF")))

    base_revenue_col = "B"  # holds base-year (year 0) revenue, feeds Year 1's growth formula
    _label(ws_fcf, "A1", "FCF Projection")
    ws_fcf[f"{base_revenue_col}3"] = "Base Year"
    ws_fcf[f"{base_revenue_col}3"].font = Font(italic=True)
    _formula_cell(ws_fcf, f"{base_revenue_col}4", "=Assumptions!B6", NUMBER_FORMAT)  # base revenue, row 4 = "Revenue"

    for t in range(1, n_years + 1):
        col = get_column_letter(t + 2)  # C, D, E, ... (B reserved for base year)
        prev_col = get_column_letter(t + 1) if t > 1 else base_revenue_col
        assumption_row = payload.assumptions.iloc[t - 1]

        ws_fcf[f"{col}2"] = t  # Year

        _input_cell(ws_fcf, f"{col}3", assumption_row["revenue_growth_pct"] / 100.0, PCT_FORMAT)
        _formula_cell(ws_fcf, f"{col}4", f"={prev_col}4*(1+{col}3)", NUMBER_FORMAT)

        _input_cell(ws_fcf, f"{col}5", assumption_row["ebit_margin_pct"] / 100.0, PCT_FORMAT)
        _formula_cell(ws_fcf, f"{col}6", f"={col}4*{col}5", NUMBER_FORMAT)

        _input_cell(ws_fcf, f"{col}7", assumption_row["da_pct_of_revenue"] / 100.0, PCT_FORMAT)
        _formula_cell(ws_fcf, f"{col}8", f"={col}4*{col}7", NUMBER_FORMAT)

        _input_cell(ws_fcf, f"{col}9", assumption_row["capex_pct_of_revenue"] / 100.0, PCT_FORMAT)
        _formula_cell(ws_fcf, f"{col}10", f"={col}4*{col}9", NUMBER_FORMAT)

        _input_cell(ws_fcf, f"{col}11", assumption_row["nwc_pct_of_revenue"] / 100.0, PCT_FORMAT)
        _formula_cell(ws_fcf, f"{col}12", f"={col}4*{col}11", NUMBER_FORMAT)

        _formula_cell(ws_fcf, f"{col}13", f"={col}6+{col}8", NUMBER_FORMAT)  # EBITDA
        _formula_cell(ws_fcf, f"{col}14", f"={col}6*(1-WACC!$B$9)+{col}8-{col}10-{col}12", NUMBER_FORMAT)  # FCF
        _formula_cell(ws_fcf, f"{col}15", f"=1/(1+WACC!$B$16)^{col}2", "0.0000")  # discount factor
        _formula_cell(ws_fcf, f"{col}16", f"={col}14*{col}15", NUMBER_FORMAT)  # PV of FCF
        ws_fcf[f"{col}16"].font = Font(bold=True)

        ws_fcf.column_dimensions[col].width = 13

    last_col = get_column_letter(n_years + 2)
    _label(ws_fcf, "A18", "Sum of PV(FCF)", bold=True)
    first_fcf_col = get_column_letter(3)
    _formula_cell(ws_fcf, "B18", f"=SUM({first_fcf_col}16:{last_col}16)", NUMBER_FORMAT)
    ws_fcf["B18"].font = Font(bold=True)

    # ------------------------------------------------------------------
    # Terminal Value
    # ------------------------------------------------------------------
    ws_tv["A1"] = "Terminal Value"
    ws_tv["A1"].font = HEADER_FONT
    ws_tv.column_dimensions["A"].width = 42
    ws_tv.column_dimensions["B"].width = 16
    ws_tv.column_dimensions["C"].width = 50

    final_fcf_cell = f"'FCF Projection'!{last_col}14"
    final_ebitda_cell = f"'FCF Projection'!{last_col}13"

    _label(ws_tv, "A3", "Terminal Growth Rate")
    _input_cell(ws_tv, "B3", payload.terminal_growth, PCT_FORMAT)
    _label(ws_tv, "A4", "Exit Multiple (x EBITDA)")
    _input_cell(ws_tv, "B4", payload.exit_multiple, MULTIPLE_FORMAT)
    _label(ws_tv, "A5", "Perpetuity Weight (0-100%)")
    _input_cell(ws_tv, "B5", payload.perpetuity_weight, PCT_FORMAT)

    _label(ws_tv, "A7", "TV (Perpetuity) = FCF_N x (1+g) / (WACC - g)", bold=True)
    _formula_cell(ws_tv, "B7", f"=IF(WACC!$B$16>B3,{final_fcf_cell}*(1+B3)/(WACC!$B$16-B3),NA())", NUMBER_FORMAT)

    _label(ws_tv, "A8", "TV (Exit Multiple) = EBITDA_N x Multiple", bold=True)
    _formula_cell(ws_tv, "B8", f"={final_ebitda_cell}*B4", NUMBER_FORMAT)

    _label(ws_tv, "A9", "Blended TV = Weight x Perpetuity + (1-Weight) x Exit", bold=True)
    _formula_cell(ws_tv, "B9", "=B5*B7+(1-B5)*B8", NUMBER_FORMAT)

    _label(ws_tv, "A10", f"PV of Terminal Value = Blended TV / (1+WACC)^{n_years}", bold=True)
    _formula_cell(ws_tv, "B10", f"=B9/(1+WACC!$B$16)^{n_years}", NUMBER_FORMAT)
    ws_tv["B10"].font = Font(bold=True)

    _label(ws_tv, "A12", "Cross-checks (implied assumptions)", bold=True)
    _label(ws_tv, "A13", "Implied Exit Multiple from Perpetuity TV (B7 / EBITDA_N)")
    _formula_cell(ws_tv, "B13", f"=IF({final_ebitda_cell}<>0,B7/{final_ebitda_cell},NA())", MULTIPLE_FORMAT)
    _label(ws_tv, "A14", "Implied Perpetuity Growth from Exit TV (solve g)")
    _formula_cell(ws_tv, "B14", f"=IF((B8+{final_fcf_cell})<>0,(B8*WACC!$B$16-{final_fcf_cell})/(B8+{final_fcf_cell}),NA())", PCT_FORMAT)

    # ------------------------------------------------------------------
    # Valuation Summary (bridge)
    # ------------------------------------------------------------------
    ws_sum["A1"] = "Valuation Summary"
    ws_sum["A1"].font = HEADER_FONT
    ws_sum.column_dimensions["A"].width = 32
    ws_sum.column_dimensions["B"].width = 16

    _label(ws_sum, "A3", "Sum of PV(FCF)")
    _formula_cell(ws_sum, "B3", "='FCF Projection'!B18", NUMBER_FORMAT)
    _label(ws_sum, "A4", "PV of Terminal Value")
    _formula_cell(ws_sum, "B4", "='Terminal Value'!B10", NUMBER_FORMAT)
    _label(ws_sum, "A5", "Enterprise Value", bold=True)
    _formula_cell(ws_sum, "B5", "=B3+B4", NUMBER_FORMAT)

    _label(ws_sum, "A6", "Net Debt (= Total Debt - Cash, or override)")
    _formula_cell(ws_sum, "B6", "=IF(Assumptions!B11<>\"\",Assumptions!B11,Assumptions!B8-Assumptions!B9)", NUMBER_FORMAT)
    _label(ws_sum, "A7", "Equity Value", bold=True)
    _formula_cell(ws_sum, "B7", "=B5-B6", NUMBER_FORMAT)

    _label(ws_sum, "A8", "Shares Outstanding")
    _formula_cell(ws_sum, "B8", "=Assumptions!B10", NUMBER_FORMAT)
    _label(ws_sum, "A9", "Value per Share", bold=True)
    _formula_cell(ws_sum, "B9", "=IF(B8>0,B7/B8,NA())", CURRENCY_FORMAT)
    ws_sum["B9"].font = Font(bold=True, size=12)

    _label(ws_sum, "A10", "Current Price")
    _formula_cell(ws_sum, "B10", "=Assumptions!B3", CURRENCY_FORMAT)
    _label(ws_sum, "A11", "Upside / (Downside) = Value/Price - 1", bold=True)
    _formula_cell(ws_sum, "B11", "=IF(B10>0,B9/B10-1,NA())", PCT_FORMAT)

    _label(ws_sum, "A13", "Note: review assumptions if upside/downside magnitude exceeds roughly 100%.")
    ws_sum["A13"].font = Font(italic=True, size=9, color="666666")
    _label(ws_sum, "A14", "Educational tool, not financial advice.")
    ws_sum["A14"].font = Font(italic=True, size=9, color="666666")

    return wb


def export_to_bytes(payload: ExportPayload) -> bytes:
    """Builds the workbook and returns it as in-memory .xlsx bytes, ready for a Streamlit download button."""
    import io
    wb = build_workbook(payload)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
